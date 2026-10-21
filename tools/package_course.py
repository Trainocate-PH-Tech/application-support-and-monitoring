"""Build instructor master and learner distributions with explicit audience boundaries."""
from pathlib import Path
import hashlib
import io
import re
import zipfile
from pptx import Presentation

ROOT=Path(__file__).resolve().parents[1]
STUDENT_FILES={
    'shared/learner-start.md','shared/assessment-rubric.md','shared/storage-worksheet.md',
    'shared/SOURCES.md','shared/linux-command-reference.md',
    'shared/sql/mysql/diagnostics.sql','shared/sql/sqlserver/diagnostics.sql',
    'shared/sql/sqlserver/workload.sql','shared/sql/sqlserver/blocking.sql',
}
STUDENT_SCRIPTS={'preflight.sh','traffic.sh','health-check.sh','collect-evidence.sh','monitor-log.sh',
                 'monitor-pipeline.sh','db-health.sh','retention.sh','README.md'}

def master_files():
    files=[ROOT/n for n in ['README.md','COVERAGE.md','VALIDATION.md','Application Support and Monitoring Course Outline 3.pdf']]
    for name in ['module1','module2','module3','shared','tools']:
        for p in (ROOT/name).rglob('*'):
            if not p.is_file() or any(x in p.parts for x in ['target','__pycache__']): continue
            if p.suffix.lower() in {'.pem','.key','.pyc'} or p.name=='.env': continue
            files.append(p)
    return sorted(files)

def learner_allowed(name):
    p=Path(name)
    if name in STUDENT_FILES: return True
    if p.parts[0] in ('module1','module2','module3'):
        return p.parent.name=='exercises' or (len(p.parts)==2 and (p.name=='README.md' or p.suffix in ('.pptx','.pdf')))
    if p.parts[:2] in [('shared','fixtures'),('shared','templates'),('shared','branding')]: return True
    return p.parent.as_posix()=='shared/scripts' and p.name in STUDENT_SCRIPTS

def student_payload(p):
    name=p.relative_to(ROOT).as_posix()
    if p.suffix=='.pptx':
        prs=Presentation(p)
        for slide in prs.slides:
            slide.notes_slide.notes_text_frame.text='Learner edition. Follow the numbered exercise and record your own evidence.'
        out=io.BytesIO();prs.save(out);return out.getvalue()
    if name=='shared/scripts/README.md':
        return ('\n'.join(line for line in p.read_text().splitlines() if not any(x in line for x in ('setup-ubuntu.sh','fault.sh','storage-lab.sh')))+'\n').encode()
    return p.read_bytes()

def build():
    files=master_files()
    for audience,filename in [('instructor','Application-Support-Training.zip'),('student','Application-Support-Student.zip')]:
        archive=ROOT/filename
        with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED) as z:
            for p in files:
                name=p.relative_to(ROOT).as_posix()
                if audience=='instructor': z.write(p,name)
                elif learner_allowed(name):
                    info=zipfile.ZipInfo.from_file(p,arcname=name)
                    info.compress_type=zipfile.ZIP_DEFLATED
                    z.writestr(info,student_payload(p))
            if audience=='student':
                z.writestr('README.md','# Application Support and Monitoring — Student edition\n\nStart with [learner setup](shared/learner-start.md), then [Module 1](module1/README.md), [Module 2](module2/README.md) and [Module 3](module3/README.md). Read the [assessment rubric](shared/assessment-rubric.md).\n\nEach day has 370 minutes of scheduled activities plus 20 minutes for checkpoints. Advanced extensions are optional. Use individual credentials and keep your work under ~/support-lab. Instructor demonstrations require the scheduled shared-stack slot. Instructor-only commands and concealed answers are supplied separately to your trainer.\n\nPowerPoint notes in this edition contain no instructor solutions. Reference scripts support practice; final assessments require your own reasoning and artifacts.\n')
        with zipfile.ZipFile(archive) as z:
            assert z.testzip() is None
            assert not any(n.endswith(('.pem','.key')) for n in z.namelist())
        print(f'{filename}: {archive.stat().st_size:,} bytes; SHA256 {hashlib.sha256(archive.read_bytes()).hexdigest()}')

if __name__=='__main__': build()
