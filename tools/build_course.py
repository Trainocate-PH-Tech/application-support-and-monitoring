"""Rebuild editable course documents and PowerPoint decks from reviewed source content."""
from pathlib import Path
import json, re
from guided_exercises import GUIDES, render_exercise
ROOT=Path(__file__).resolve().parents[1]
MODULES={1:[],2:[],3:[]}
def lab(m,slug,title,minutes,scenario,concepts,steps,commands,checks,tricks,answer,output):
    MODULES[m].append(dict(slug=slug,title=title,minutes=minutes,scenario=scenario,concepts=concepts,steps=steps,commands=commands,checks=checks,tricks=tricks,answer=answer,output=output))

lab(1,'01-access-users','Access, identity and user creation',40,
'You are the new support team. Establish individual identities before changing the shared loyalty server.',
['The SSH key authenticates a login; it does not make the login root.','A user, primary group, supplementary groups and home-directory permissions are separate controls.','Use sudo only for the supervised administrative step; keep routine investigation unprivileged.'],
['Instructor verifies the supplied key locally and demonstrates the README connection. Learners receive their own account/key; never copy the instructor key.','Record whoami, id, hostname and UTC time. Explain the meaning of ~, /, ., .. and the shell prompt.','Instructor substitutes a unique learner name below, creates the account, then installs that learner’s PUBLIC key. Do not grant all learners unrestricted sudo.','Reconnect as the learner, create ~/support-lab/evidence, and prove another learner cannot read a private file. Create a shared team group with the instructor.'],
'''# Instructor workstation; keep private key here
chmod 600 okada.pem
ssh -i okada.pem ubuntu@47.131.221.162
# On server
whoami; id; hostname; date -u +%FT%TZ
# Instructor-only example; change learner01/team01 for each team
sudo adduser --disabled-password --gecos '' learner01
sudo groupadd -f team01
sudo usermod -aG team01 learner01
sudo chmod 750 /home/learner01
sudo install -d -m 700 -o learner01 -g learner01 /home/learner01/.ssh
# Put only learner01's PUBLIC key in authorized_keys, then:
sudo chown learner01:learner01 /home/learner01/.ssh/authorized_keys
sudo chmod 600 /home/learner01/.ssh/authorized_keys
# As learner after reconnecting
mkdir -p ~/support-lab/evidence
umask 077
printf 'private lab note\\n' > ~/support-lab/private.txt
id''',
['Learner logs in using their own key and has a writable home.','Private file has mode 600; team membership appears after a new login.'],
['Does a # character inside a command prove you are root?','Why might id in an old shell omit a newly granted group?','Does chmod 777 fix a missing filesystem quota?'],
'Use id -u to establish identity. A new login refreshes supplementary groups. Directory traversal requires execute permission; read permission alone is insufficient. chmod changes access, not capacity. Keep ~/.ssh owned by its account and never place a private key in the course repository.',
'Identity baseline, permission evidence and a short access checklist.')

lab(1,'02-files-packages','Navigation, files and package installation',65,
'A deployment arrives as an archive. Find the right directory, inspect it, copy it safely and install diagnostic tools.',
['Absolute paths start at /; relative paths start at the current working directory.','Quote paths and distinguish copying a file from overwriting an existing destination.','apt is the Ubuntu package manager; dnf examples are Red Hat comparisons, not commands for this host.'],
['Build a directory tree under ~/support-lab. Navigate using absolute and relative paths, then recover using cd -.','Create, copy, rename and remove disposable files. Inspect ownership, symlinks, hidden files, PATH and command lookup.','List the supplied fixture archive before extraction into a new directory. Explain why archives from unknown sources need review.','Instructor demonstrates apt update, package search, installation and package ownership. Students verify curl, nc, dig, lsof and pidstat.'],
'''mkdir -p ~/support-lab/{incoming,work,archive}
cd ~/support-lab
pwd; ls -lah
printf 'alpha\\nbeta\\n' > 'incoming/status report.txt'
cp -i -- 'incoming/status report.txt' work/
mv -- 'work/status report.txt' work/status.txt
head -n 1 work/status.txt
ln -s status.txt work/latest
ls -l work; readlink work/latest
cd /var/log; pwd; cd -
printf '%s\\n' "$PATH"
command -v curl
# Instructor; package installation changes the shared host
sudo apt-get update
sudo apt-get install -y curl netcat-openbsd dnsutils lsof sysstat
# Fixture paths assume you are back at the course repository root
tar -tzf shared/fixtures/support-fixtures.tar.gz
mkdir -p ~/support-lab/fixtures
tar -xzf shared/fixtures/support-fixtures.tar.gz -C ~/support-lab/fixtures
# Target only a file you created
rm -i -- ~/support-lab/work/status.txt''',
['Student explains why the symlink breaks after its target is removed.','Commands are found through PATH; package installation is verified rather than assumed.'],
['Does cd /var/log change another terminal’s working directory?','Why is sudo echo text > /root/file different from echo text | sudo tee /root/file?','What happens when an unquoted variable contains spaces?'],
'The shell performs redirection before sudo runs echo. Use a deliberately approved destination and sudo tee when needed. cd changes the current process environment. A symlink stores a path and does not keep its target alive. Use -- before filenames when the tool supports it.',
'A directory tree, package verification table and safe file-handling notes.')

lab(1,'03-storage','Allocate space and diagnose disk limits',50,
'A learner reports “my EC2 user is capped.” Determine which layer is limiting writes before selecting a fix.',
['EBS volume → partition → filesystem → directory/account; growing one layer does not automatically grow all layers.','Compare byte capacity, inode capacity, quotas, permissions and deleted-open files.','The observed host already has an expanded root filesystem; reproduce growth on a disposable disk.'],
['Capture lsblk, findmnt, df -hT, df -i and quota evidence. Interpret unavailable quota tools separately from proof of no quota.','Instructor runs the disposable storage helper after installing parted. Record capacity before creation, after device growth and after filesystem growth. Read the helper before executing it.','Use the bounded disk-pressure and quota procedures in shared/instructor/storage.md. All writes stay on the disposable image.','Explain the EBS expansion workflow and ext4 versus XFS tools. Never copy device names from an example into a production command.'],
'''lsblk -o NAME,SIZE,TYPE,FSTYPE,MOUNTPOINTS
findmnt -no SOURCE,FSTYPE,OPTIONS /
df -hT; df -i
quota -s || true
sudo repquota -a
sudo lsof +L1
# Instructor on disposable lab host
sudo apt-get install -y parted quota cloud-guest-utils
sudo shared/scripts/storage-lab.sh create
sudo lsblk; sudo df -hT /var/lib/support-storage-lab/mount
sudo shared/scripts/storage-lab.sh grow
sudo shared/scripts/storage-lab.sh cleanup
# Real EBS workflow is documented, not executed against the root disk.''',
['Student identifies the limiting layer from evidence, not the wording of the complaint.','Disposable filesystem grows and can be cleaned up; root partition is untouched.'],
['df says full but du is small: name three explanations.','Will adding a user allocate a new partition?','Why can root sometimes write when an ordinary user cannot, even without quotas?'],
'Possible causes include deleted-open files, hidden data under mounts, inaccessible paths during du, reserved ext4 blocks, inode exhaustion, quotas, or permissions. Normal user creation does not create a per-user EBS disk. Show inode and quota evidence before diagnosing a cap. EBS growth requires appropriate AWS access and a snapshot; grow the partition only if present, then the matching filesystem.',
'A storage decision tree and before/after capacity evidence.')

lab(1,'04-command-investigations','Command pipelines and log investigations',65,
'Operations received mixed logs and filenames. Produce accurate incident counts without losing records or matching misleading text.',
['grep selects lines; awk interprets fields; sed transforms text; each has different quoting rules.','Sort before uniq, choose numeric sorting for numbers and preserve filenames with NUL separators.','Counts require a clear unit: lines, requests, users and incidents are not interchangeable.'],
['Use the extracted fixtures to count 5xx requests and rank paths by frequency. Compare a naive ERROR search against parsed status codes.','Find files with spaces and leading dashes. Archive selected files using find -print0 and tar --null.','Extract request IDs, inspect head/tail, follow a log and stop following without killing the application.','Complete the command challenge bank. Explain each pipeline stage and capture exit codes with and without pipefail.'],
'''cd ~/support-lab/fixtures
less application.log
head -n 4 application.log; tail -n 4 application.log
grep -n 'status=5[0-9][0-9]' application.log
awk '{for(i=1;i<=NF;i++) if($i ~ /^path=/) print $i}' application.log | sort | uniq -c | sort -nr
cut -d, -f2 latency.csv | tail -n +2 | sort -n
sed -n '2,6p' application.log
find filenames -type f -print0 | xargs -0 -r wc -c
find filenames -type f -print0 | tar --null -T - -czf ~/support-lab/archive/selected.tar.gz
printf 'seen\\n' | tee ~/support-lab/evidence/pipeline.txt
wc -l application.log
set -o pipefail
grep 'NO_SUCH_REQUEST' application.log | wc -l
printf 'pipeline_exit=%s\\n' "$?"
date -u +%FT%TZ''',
['Counts match parsed fixture records; malformed rows are reported separately.','All unusual filenames survive the archive round trip.'],
['Does grep returning 1 always mean an error?','Will uniq find nonadjacent duplicates?','Can cut reliably parse arbitrary CSV with quoted commas?'],
'grep exits 1 for no matches, 2 for an error. uniq removes adjacent duplicates only. Basic cut is not a full CSV parser. A pipeline usually reports its last command status; pipefail changes this. Use NUL separators for arbitrary filenames. The fixture contains six valid requests, two 5xx errors, two slow requests at >=1000 ms, and one malformed row.',
'Annotated pipelines, result counts and an archive listing.')

lab(1,'05-install-stack','Install and inspect the application stack',60,
'Deploy the loyalty application, establish a healthy baseline and understand the service before introducing failures.',
['A Java runtime, servlet container, database and application configuration are separate dependencies.','systemd unit configuration defines process identity, environment, startup and restart behavior.','A live process is weaker evidence than an authenticated business request.'],
['Instructor reads and runs setup-ubuntu.sh on a prepared training host. Learners explain apt, archive checksum, Maven build and WAR deployment stages.','Inspect the service unit and database/application logs. Identify the service user, Java memory cap and restart policy.','Test liveness, readiness, player lookup and report routes. Record status, response body and request ID.','Explain each layer when /live succeeds but /health fails. Use SSH forwarding for browser access.'],
'''# Instructor only; installation must be rehearsed before class
sudo shared/scripts/setup-ubuntu.sh --training-host
systemctl cat support-tomcat
systemctl status support-tomcat --no-pager
journalctl -u support-tomcat --since '10 minutes ago' --utc --no-pager
curl -i http://127.0.0.1:8080/loyalty/live
curl -i http://127.0.0.1:8080/loyalty/health
curl -i 'http://127.0.0.1:8080/loyalty/player?id=1'
ss -lnt
# Learner workstation, using their own key/account
ssh -L 8080:127.0.0.1:8080 learner01@47.131.221.162
# Then browse http://127.0.0.1:8080/loyalty/health locally.''',
['Readiness returns 200 and player 1 returns synthetic data.','Student explains why Tomcat runs as supportapp and MySQL is not exposed publicly.'],
['Does systemctl enable start a stopped service?','Can /live return 200 while the database password is wrong?','Why should you collect the journal before restart?'],
'enable configures boot behavior; --now also starts. /live intentionally avoids the database. /health opens a connection and executes SELECT 1. Restart can destroy transient evidence and temporarily hide a fault. systemctl cat reveals drop-ins as well as the main unit.',
'A component diagram and a healthy baseline evidence bundle.')

lab(1,'06-service-log-triage','Service failures and application log diagnosis',35,
'The application worked before a change. One team investigates a startup fault; another receives a permission fault after reset.',
['Capture unit state, exit status and journal before changing configuration.','Correlate an HTTP response request ID with its application log line.','Distinguish startup failure, runtime failure and failed dependency checks.'],
['Instructor injects startup; learners record systemctl status, journal and effective unit. Identify the invalid executable without guessing.','Instructor resets, then injects permissions. Compare /live and /health with the startup symptom.','Find the failing path and ownership using namei and stat under instructor supervision. Propose the narrow permission repair.','Reset and collect fresh successful requests; explain the difference between mitigation and recovery verification.'],
'''# Instructor: one fault at a time
sudo shared/scripts/fault.sh startup
systemctl status support-tomcat --no-pager
systemctl cat support-tomcat
journalctl -u support-tomcat -n 60 --utc --no-pager
sudo shared/scripts/fault.sh reset
sudo shared/scripts/fault.sh permissions
curl -i http://127.0.0.1:8080/loyalty/health
sudo namei -l /etc/support-lab/app.properties
sudo stat /etc/support-lab/app.properties
sudo shared/scripts/fault.sh reset''',
['Evidence names the actual failed executable or unreadable file.','Recovery includes HTTP and business checks, not only active service status.'],
['Does a restart policy repair a bad executable path?','Does a missing ERROR line prove the service started?','Why is chmod 777 an inappropriate application-config fix?'],
'The startup drop-in overrides ExecStart; inspect both unit and drop-in. The permissions fault leaves root owning a mode-600 config, so supportapp cannot read it. Correct access is root:supportapp mode 640. Repeated restarts cannot fix either cause. Preserve the failure evidence before using reset.',
'A timeline with failing and recovered request IDs.')

lab(1,'07-resources-network','Resource and network troubleshooting',30,
'Users report intermittent errors. Determine whether CPU, memory, DNS, TCP, TLS or database authentication evidence supports the complaint.',
['High load is not the same measurement as CPU utilization; memory available differs from free.','A listening port and successful TCP connection do not prove database authentication.','Ubuntu firewall checks differ from Red Hat SELinux indicators; do not disable controls as a diagnostic shortcut.'],
['Capture interval CPU/process data, memory available, disk and file-handle evidence. Use the bounded CPU demonstration in the instructor guide.','Test the layers in order: resolution, route indicators, listener, TCP, HTTP/TLS, application/database authentication.','Instructor injects db-auth. Compare nc success with failing /health and inspect the request log.','Analyze the DNS/TLS/SELinux evidence cards and propose the next discriminating check.'],
'''uptime; free -h
ps -eo pid,ppid,user,comm,pcpu,pmem --sort=-pcpu | head
top -b -n 1 | head -n 20
pidstat 1 3
sudo lsof -p "$(systemctl show -p MainPID --value support-tomcat)" | head
ss -lnt; nc -vz -w 2 127.0.0.1 3306
getent hosts localhost; nslookup localhost; dig localhost
ping -c 2 127.0.0.1; traceroute -m 3 127.0.0.1
curl --max-time 5 -v http://127.0.0.1:8080/loyalty/health
sudo shared/scripts/fault.sh db-auth
curl -i http://127.0.0.1:8080/loyalty/health
sudo shared/scripts/fault.sh reset''',
['Student differentiates TCP reachability from successful authenticated queries.','Diagnosis cites interval evidence and acknowledges ICMP/filtering limitations.'],
['Does failed ping prove the host is down?','Is low free memory alone proof of memory pressure?','Does curl -k repair an expired certificate?'],
'ICMP can be blocked while TCP works. Inspect available memory, swap/OOM evidence and workload behavior. curl -k bypasses validation and does not repair certificates. A MySQL listener can accept TCP while rejecting credentials. Red Hat AVC evidence requires policy/path investigation; do not teach disabling SELinux as the fix.',
'A layered connectivity table with evidence and remaining unknowns.')

lab(1,'08-first-script-capstone','First script and Linux incident challenge',45,
'A deployment left the loyalty application unavailable. Your team has 25 minutes to identify the fault and deliver a reusable evidence collector.',
['Build a small script from understood commands, then test failure behavior.','Record command failures without losing the rest of the incident evidence.','A recovery claim needs a customer-facing check and a clear incident summary.'],
['Write a Bash script with a shebang, timestamp, service state, disk use and recent errors. Run bash -n, make it executable and run it from another directory.','Instructor selects startup, db-auth or permissions without showing the selection. Assign investigator, scribe and reviewer.','Collect evidence before repair, state a hypothesis, request the instructor-controlled mitigation and verify recovery.','Compare your script with the reference collector; identify one missing timeout, permission or redaction control.'],
'''mkdir -p ~/support-lab/bin
# Write ~/support-lab/bin/triage.sh using your editor
bash -n ~/support-lab/bin/triage.sh
chmod 750 ~/support-lab/bin/triage.sh
~/support-lab/bin/triage.sh
# Reference implementation; run from the course repository root
shared/scripts/collect-evidence.sh ~/support-lab/evidence '15 minutes ago'
printf 'collector_exit=%s\\n' "$?"
find ~/support-lab/evidence -maxdepth 1 -name '*.tar.gz' -print''',
['Script runs without interactive input and records UTC timestamps.','Incident report includes evidence, probable cause, mitigation and recovery validation.'],
['Why can an evidence collector legitimately exit nonzero and still produce a useful bundle?','Does chmod +x guarantee a script will run on every shell?'],
'A partial bundle may contain a failed service-status check or a permission-denied log capture. Record those statuses instead of suppressing them. The interpreter in the shebang must exist; use Bash explicitly for Bash syntax. Grade reasoning and verified recovery above speed.',
'Executable triage script, checklist, evidence bundle and incident report.')

lab(1,'09-command-challenge-bank','Advanced Linux command challenge bank','Extension',
'You must produce defensible answers from awkward input without changing source evidence.',
['Predict command behavior before running it.','Preserve evidence and explain assumptions about input format.','Robust pipelines handle empty input, malformed records and unusual filenames.'],
['Solve all 18 questions in shared/fixtures/command-challenges.md. Save commands, output and explanations.','Create one counterexample that breaks a naive pipeline, then repair it.','Explain which commands are GNU/Linux-specific and what you would check on a different host.'],
'''# Examples to investigate in a disposable directory
printf '2\\n10\\n1\\n' | sort
printf '2\\n10\\n1\\n' | sort -n
printf 'a\\nb\\na\\n' | uniq -c
printf 'a\\nb\\na\\n' | sort | uniq -c
printf 'one\\ntwo' | wc -l
find ~/support-lab/fixtures/filenames -type f -print0 | xargs -0 -r ls -ld --''',
['Each answer includes an explanation and a counterexample where appropriate.','No command mutates original evidence or uses blanket deletion.'],
['Is line count always record count?','When does tail -f miss a replacement log file?','Can grep ERROR count HTTP failures accurately?'],
'Use tail -F for following a filename across rotation when appropriate. wc -l counts newline characters. Parse explicit status fields rather than unstructured words. See instructor command-answer key for all 18 explanations.',
'A reviewed command notebook; extensions do not displace required class activities.')

lab(2,'01-db-baseline','Database baseline and safe access',35,
'The application reports database errors. Establish engine, endpoint, identity and connectivity before tuning anything.',
['Use a protected client option file instead of putting passwords in command arguments.','Read-only application rights differ from administrative diagnostic permissions.','Capture UTC time, engine version and connection counts with every baseline.'],
['Instructor provides a training login with the documented diagnostic grants. Learners create a mode-600 MySQL client file.','Run db-health.sh and diagnostics.sql; label permission-denied results rather than treating them as empty findings.','Compare healthy HTTP, TCP and authenticated SQL results. Estimate connection headroom from observed counts and max_connections.'],
'''# Edit ~/.support-lab.cnf: [client], user, password, host=127.0.0.1
chmod 600 ~/.support-lab.cnf
MYSQL_CNF="$HOME/.support-lab.cnf" shared/scripts/db-health.sh
mysql --defaults-extra-file="$HOME/.support-lab.cnf" < shared/sql/mysql/diagnostics.sql
# Interactive mysql examples
# SELECT UTC_TIMESTAMP(),CURRENT_USER(),VERSION();
# SHOW FULL PROCESSLIST;
# SHOW VARIABLES LIKE 'max_connections';''',
['Baseline identifies endpoint and effective database user.','Learner distinguishes access denied from server unreachable.'],
['Can a successful SELECT 1 prove every business query is fast?','Does SHOW PROCESSLIST necessarily show all users to an unprivileged login?'],
'SELECT 1 checks a narrow path. PROCESS privilege controls visibility of other sessions; missing privileges can hide incident evidence. Do not distribute the application account password to learners. Use per-team diagnostic accounts and instructor-controlled writes.',
'A database baseline and a protected client configuration, never submitted with its password.')

lab(2,'02-mysql-slow','MySQL slow-query capture and ranking',45,
'Player reports are slow. Capture an intentional slow query and rank work by both frequency and total time.',
['Slow logging is a configuration change and must be restored after the exercise.','Global long_query_time becomes the default for new sessions; existing sessions may retain their value.','Frequency, total time, rows examined and lock time answer different questions.'],
['Instructor records existing slow-log settings, enables logging briefly and restores them after capture. Use a fresh connection for the test.','Run a bounded SLEEP query and report query. Locate the configured log file without assuming its path.','Summarize the supplied slow log with mysqldumpslow using count and time ordering. Compare the rankings.'],
'''# Instructor mysql session; record values BEFORE changing them
# SHOW VARIABLES WHERE Variable_name IN ('slow_query_log','long_query_time','slow_query_log_file');
# SET GLOBAL slow_query_log=ON;
# SET GLOBAL long_query_time=0.2;
# Fresh training connection:
# SET SESSION long_query_time=0.2;
# SELECT SLEEP(0.3);
# SELECT COUNT(*),SUM(amount) FROM support_lab.transactions WHERE player_id=42;
mysqldumpslow -s t -t 5 shared/fixtures/mysql-slow.log
mysqldumpslow -s c -t 5 shared/fixtures/mysql-slow.log
# Instructor restores the exact original global settings, not guessed defaults.''',
['Slow record appears with query time, lock time and row counts.','Report explains why a frequent moderate query may cost more overall than one rare slow query.'],
['Does Query_time alone prove CPU pressure?','Will changing a global threshold rewrite existing sessions?'],
'Elapsed time includes waiting and is not a CPU measurement. Use fresh sessions or an explicit session setting. Supplied evidence has a frequent player-report query and a rarer long query; compare total impact rather than selecting only the largest individual duration.',
'A ranked slow-query table with units, frequency and next checks.')

lab(2,'03-mysql-plans','EXPLAIN, indexes and misleading assumptions',35,
'The report scans many transactions. Decide whether a targeted index improves this workload.',
['EXPLAIN estimates access strategy and cardinality; it is not proof of actual elapsed time.','Indexes add write/storage costs and must be justified by workload.','A tiny training dataset may show plan differences without dramatic timing differences.'],
['Capture EXPLAIN and a baseline report result for player 42. Record data volume.','Instructor creates the training index; capture the new plan and run the same query again.','Compare access type, selected key and estimated rows. Keep business results unchanged. Drop the index to restore the exercise baseline.'],
'''# Run in an interactive instructor-approved mysql session
# USE support_lab;
# SELECT COUNT(*) FROM transactions;
# EXPLAIN SELECT COUNT(*),SUM(amount) FROM transactions WHERE player_id=42;
# CREATE INDEX idx_transactions_player ON transactions(player_id);
# EXPLAIN SELECT COUNT(*),SUM(amount) FROM transactions WHERE player_id=42;
# SELECT COUNT(*),SUM(amount) FROM transactions WHERE player_id=42;
# DROP INDEX idx_transactions_player ON transactions;
# Do not use EXPLAIN ANALYZE on an unknown expensive query: it executes it.''',
['Before/after plans and equivalent query results are retained.','Recommendation includes an index tradeoff and measurement limitation.'],
['Is every full scan bad?','Can an index be present yet unused?','Does LIMIT guarantee cheap execution?'],
'Full scans may be optimal for small tables or low-selectivity predicates. Expressions, types, selectivity and statistics affect index use. LIMIT can still require scanning or sorting many rows. For 100,000 seeded rows, player 42 has 1,000 transactions totaling 10,000.00.',
'A before/after plan comparison and an evidence-based recommendation.')

lab(2,'04-mysql-locks','MySQL blocking and deadlocks',40,
'Two concurrent updates stall. Separate an ordinary lock wait from a deadlock and recover both sessions.',
['A blocking chain can persist without a cycle; a deadlock requires a cycle.','The deadlock victim is not necessarily the root cause.','An open transaction can hold locks even while its client appears idle.'],
['Open two writer sessions and one observer. Follow contention.sql to create a simple blocking case first.','Reset both sessions, then lock players in opposite order to create a deadlock.','Capture InnoDB status promptly; identify held/requested locks, victim and survivor. Roll back both sessions.','Propose consistent lock ordering, short transactions and bounded retry behavior.'],
'''cat shared/sql/mysql/contention.sql
# Observer while sessions are blocked:
# SELECT * FROM performance_schema.data_lock_waits;
# SHOW FULL PROCESSLIST;
# SHOW ENGINE INNODB STATUS\\G
# Recovery in BOTH writer sessions:
# ROLLBACK;''',
['Evidence shows the blocking relation or deadlock cycle.','Both sessions end their transactions and a subsequent update succeeds.'],
['Does killing the waiting session remove the head blocker?','Does retrying immediately forever solve deadlocks safely?'],
'Killing a waiter does not release locks held by the blocker. Deadlock detection rolls back a victim; the survivor may still need an explicit rollback in this exercise. Retry only a complete safe transaction, with bounds/backoff and idempotency considerations.',
'A lock diagram, deadlock interpretation and verified cleanup.')

lab(2,'05-replication','Replication delay and reporting freshness',15,
'A reporting screen shows stale points while the primary accepts transactions. This is a supplied evidence exercise, not a configured replica.',
['Application routing and replica freshness must be investigated together.','Replication threads, errors and delay metrics need timestamps and context.','An empty replica-status result on a standalone server is not a replication failure.'],
['Read replication.json and correlate primary write, replica read and captured status.','Identify the SQL applier error and distinguish unknown lag from zero lag.','Propose routing or repair options with freshness/load tradeoffs; do not skip failed transactions blindly.'],
'''cat shared/fixtures/replication.json
# On a real configured replica, with suitable privilege:
# SHOW REPLICA STATUS\\G
# Compare Source/Replica fields, applier error, timestamps and routing.
# Never run RESET REPLICA or skip transactions as an unreviewed quick fix.''',
['Student recognizes NULL delay as unknown rather than healthy.','Recommendation addresses stale reads and the replication error separately.'],
['Does a running receiver thread prove the applier is current?','Can reporting freshness be inferred from one lag number alone?'],
'The synthetic replica receiver is running but SQL applier stopped with duplicate-key error 1062. NULL lag is not zero. Investigate consistency and cause before repair; temporary primary reads require capacity and business approval.',
'A replication escalation note with evidence and routing tradeoffs.')

lab(2,'06-sqlserver-baseline','SQL Server requests, sessions and wait deltas',40,
'SQL Server is blamed for high CPU. Capture active work and distinguish current activity from cumulative counters.',
['SQL Server 2022 diagnostic permissions differ from older examples; validate visibility first.','Requests and sessions are related but not interchangeable.','Wait statistics are cumulative; compare two captures over a known interval.'],
['Connect to the separate SQL Server lab using the instructor-provided client settings. Run diagnostics.sql.','Execute the bounded workload from a second connection and capture requests, cached query metrics and plans.','Use sqlserver-waits.csv to calculate wait deltas and explain CPU versus waiting. Query plans may expose sensitive SQL text: review before sharing.'],
'''# SQLCMDPASSWORD is entered locally; do not put passwords in -P arguments.
# sqlcmd -S tcp:SQL_HOST,1433 -U TEAM_LOGIN -d SupportLab -i shared/sql/sqlserver/diagnostics.sql
# Run workload.sql in a second connection while observing.
cat shared/fixtures/sqlserver-waits.csv
# SQL example:
# SELECT session_id,status,cpu_time,total_elapsed_time,wait_type,blocking_session_id
# FROM sys.dm_exec_requests WHERE session_id<>@@SPID;''',
['Report includes time window, permissions and correct units.','Student separates historical high waits from interval changes.'],
['Can a sleeping session be the head blocker?','Does highest total_worker_time identify the currently slowest query?'],
'A sleeping session may hold an uncommitted transaction. Cached totals aggregate executions since compilation and do not necessarily describe this incident. Inspect execution_count, interval evidence and plans together. In the fixture LCK_M_X grows 12,000 ms, SOS_SCHEDULER_YIELD 1,000 ms, and PAGEIOLATCH_SH 200 ms.',
'A SQL Server baseline and a wait-delta interpretation.')

lab(2,'07-sqlserver-blocking-plans-jobs','Blocking chains, Query Store and job failures',45,
'A deployment coincides with latency and a failed reporting job. Establish whether the symptoms share a cause.',
['Identify the head blocker through sessions as well as active requests.','Query Store compares plans and runtime intervals; correlation with deployment is not proof of causation.','SQL Agent job outcome and step outcome need separate interpretation.'],
['Use blocking.sql in three connections. Identify A as blocker, observe B’s timeout and roll back A.','Run workload.sql before and after the approved index change. Inspect Query Store; flush its data as instructor if needed.','Instructor starts SupportLab-IntentionalFailure; learners inspect its history and message.','Analyze query-store.csv and job-history.csv for a deterministic plan-regression narrative, clearly labeled synthetic.'],
'''# SQL Server sessions; see blocking.sql for exact sequencing
# EXEC msdb.dbo.sp_start_job @job_name=N'SupportLab-IntentionalFailure';
# USE SupportLab;
# EXEC sys.sp_query_store_flush_db; -- instructor-controlled
# SELECT * FROM sys.database_query_store_options;
cat shared/fixtures/query-store.csv
cat shared/fixtures/job-history.csv
# Cleanup: ROLLBACK in the transaction owner; drop only the training index.''',
['Head blocker is identified before mitigation.','Query Store findings name query ID, plan IDs, interval and metric units.','Job diagnosis cites the failed step and message.'],
['Does a changed plan guarantee a regression?','Is plan forcing always the safest first action?','Can a failed job explain stale reports without explaining API latency?'],
'Compare runtime evidence across comparable intervals and row counts. Plan forcing needs validation and rollback. The fixture shows query 17 changing from plan 3 to 8 with higher duration/reads after deployment, while the job failure is a separate permission issue. The live training job deliberately throws a synthetic error.',
'A blocking worksheet, plan comparison and job-failure escalation.')

lab(2,'08-app-db-correlation','Application/database correlation',55,
'An HTTP request took five seconds, but its database query took 40 milliseconds. Determine where the rest of the time went.',
['Application latency includes pool waits, SQL time, retries and downstream work.','Correlate IDs and UTC windows; acknowledge clock skew and missing instrumentation.','Request IDs in app logs do not automatically appear in database logs.'],
['Join the supplied application and database evidence conceptually using the documented operation mapping.','Build a timeline from correlation.csv and identify pool wait and repeated attempts.','Compare the live latency fault, which adds application delay, with the synthetic database-contention scenario.','Draft an escalation with facts, hypotheses and the next measurement required.'],
'''cat shared/fixtures/correlation.csv
sudo shared/scripts/fault.sh latency
shared/scripts/traffic.sh http://127.0.0.1:8080/loyalty 5
journalctl -u support-tomcat --since '5 minutes ago' --utc --no-pager
sudo shared/scripts/fault.sh reset
# The sample app opens a connection per request; pool exhaustion is evidence-based here.''',
['Timeline accounts for measured time without inventing an exact database-request-ID join.','Learner distinguishes the sample app’s behavior from the supplied pooled-app evidence.'],
['Does slow HTTP prove slow SQL?','Can retries amplify a partial outage?'],
'The supplied request waits 4,000 ms for a pool connection, spends 40 ms in SQL and 960 ms elsewhere. The live latency fault intentionally delays after connecting; it is not a real pool implementation. Repeated retries can increase load and duplicate non-idempotent work.',
'A correlated timeline and a targeted escalation request.')

lab(2,'09-database-capstone','Final database incident challenge',80,
'Teams receive a hidden live database-auth fault plus independent SQL Server and replication evidence. Separate primary cause from secondary symptoms.',
['Time-box investigation and preserve the evidence behind each conclusion.','Use read-only diagnostics first; route writes and session termination through the instructor.','Recovery requires application checks and database checks.'],
['Spend 10 minutes assigning roles and defining impact; 30 minutes investigating; 15 minutes proposing and verifying mitigation; 25 minutes reporting and review.','Deliver a ranked MySQL slow-query report and SQL Server worksheet covering waits, blockers, Query Store and jobs.','Write a database-health script outline with authentication, timeouts, timestamped output and explicit partial-failure handling.','Answer counterfactual questions: what evidence would disprove your diagnosis, and what would you monitor next?'],
'''shared/scripts/db-health.sh
shared/scripts/health-check.sh
# Copy templates into your own workspace before editing
cp shared/templates/database-report.md ~/support-lab/database-report.md
cp shared/templates/incident-report.md ~/support-lab/database-incident.md''',
['Reports distinguish live observations from synthetic evidence.','Recommended mitigation has owner, rollback and recovery criteria.'],
['Should every warning found during an incident be treated as its root cause?','Can you claim recovery from one successful request?'],
'The instructor uses db-auth for the live outage; TCP and database service remain up. Independent evidence requires independent diagnoses. Validate several bounded business requests and note the observation window. Apply the rubric to evidence and reasoning rather than one prescribed wording.',
'MySQL report, SQL Server worksheet, timeline and health-check outline.')

lab(3,'01-bash','Bash foundations for support automation',65,
'Turn yesterday’s commands into a script another engineer can run reliably.',
['Quote expansions; treat arguments as data; use arrays for command arguments.','Exit codes are an interface: document healthy, unhealthy and usage outcomes.','set -e is not a substitute for deliberate error handling.'],
['Write a script accepting a service name and URL, with functions for timestamped logging and checks.','Add a loop over checks, conditionals, command substitution and a usage branch.','Test healthy, stopped-service, invalid-argument and unreachable-endpoint cases. Compare with health-check.sh.'],
'''bash -n shared/scripts/health-check.sh
SERVICE=support-tomcat URL=http://127.0.0.1:8080/loyalty/health shared/scripts/health-check.sh
printf 'exit=%s\\n' "$?"
URL=http://127.0.0.1:1 shared/scripts/health-check.sh
printf 'exit=%s\\n' "$?"
DISK_LIMIT=invalid shared/scripts/health-check.sh
printf 'exit=%s\\n' "$?"''',
['Script uses bounded network calls and preserves failure status.','Usage errors return 2; unhealthy checks return 1; healthy checks return 0.'],
['What happens to $? after echo runs?','Why can set -e behave differently inside conditionals?'],
'Capture the status immediately after the command. Explicit if command; then branches make expected failures clear. The reference health checker validates the disk threshold and reports each check instead of stopping at the first unhealthy service.',
'A documented executable Bash health checker and test transcript.')

lab(3,'02-evidence','Automated evidence collection',60,
'An escalation needs a portable evidence bundle even when some commands fail or lack permission.',
['Use restrictive permissions, unique timestamped directories and a manifest.','Capture stdout, stderr and exit status for each check.','Redaction needs review; a regex cannot guarantee removal of every secret.'],
['Run the reference collector and inspect its manifest, archive permissions and contents.','Run once without journal access and once with the training diagnostic permissions; explain partial results.','Add a bounded time-window capture and a synthetic secret marker; confirm masking while avoiding real secrets.','Review the archive before sharing and write a short index explaining the evidence.'],
'''shared/scripts/collect-evidence.sh ~/support-lab/evidence '20 minutes ago'
rc=$?; printf 'collector=%s\\n' "$rc"
find ~/support-lab/evidence -name manifest.tsv -exec cat {} \\;
# Inspect the printed archive path; do not automatically upload it.
# tar -tzf /exact/path/printed/by/collector.tar.gz
# stat /exact/path/printed/by/collector.tar.gz''',
['Failed checks remain visible in the manifest.','Bundle contains no credentials or environment dumps and is reviewed before sharing.'],
['Can a timestamp alone prevent simultaneous filename collisions?','Should permission denied be silently converted to an empty success file?'],
'mktemp adds uniqueness; umask 077 restricts access. Nonzero capture results set partial=1 while preserving the rest of the bundle. Do not collect unrestricted environment/process argument data. Structured sensitive values may need stronger, format-aware redaction.',
'A reviewed evidence archive, manifest and evidence index.')

lab(3,'03-monitoring','Log parsing and threshold monitoring',45,
'Build a monitor that counts actual requests, rejects malformed rows and reports slow requests.',
['Define the time window, denominator, threshold comparison and no-data behavior.','Count status fields rather than words such as ERROR in free text.','A monitoring result is only as useful as its handling of missing or stale input.'],
['Run monitor-log.sh on the fixture and reconcile every output field manually.','Test a healthy-only window, an empty file and a malformed-only file.','Slice a fresh journal window before running the parser; explain why cumulative logs distort an error-rate alert.','Add a freshness check as an extension and document threshold inclusivity.'],
'''shared/scripts/monitor-log.sh shared/fixtures/application.log 10 1000
printf 'exit=%s\\n' "$?"
journalctl -u support-tomcat --since '5 minutes ago' --utc -o cat --no-pager > ~/support-lab/window.log
shared/scripts/monitor-log.sh ~/support-lab/window.log 10 1000
# Bound and inspect input; ordinary Tomcat startup lines count as malformed, not requests.''',
['Fixture result is 6 requests, 2 errors, 33.33%, 2 slow, 1 malformed, ALERT.','Empty input yields NO_DATA and exit 2, never healthy.'],
['Is zero observed errors equivalent to a healthy service when no requests exist?','Should 404 responses count as server failures in this policy?'],
'The supplied policy counts status >=500 as server failures and latency >= threshold as slow. A rate equal to the threshold alerts. NO_DATA is unknown, not green. HTTP 4xx monitoring may be added separately according to business context.',
'A monitoring CSV and a written policy with denominator and no-data semantics.')

lab(3,'04-slow-pipeline','Slow-query summaries and correlation pipelines',20,
'Create one summary combining application errors with database slow-query evidence without pretending they share a guaranteed join key.',
['Use database-native summaries for multiline slow logs.','Keep units, source window and provenance in exported results.','A correlation hypothesis must be labeled if an exact join is unavailable.'],
['Summarize the supplied MySQL slow log by frequency and time.','Combine app-monitor CSV and slow-query summary in a timestamped report directory.','Explain why grepping individual SELECT lines can lose multiline query context.'],
'''mkdir -p ~/support-lab/summary
mysqldumpslow -s t -t 5 shared/fixtures/mysql-slow.log > ~/support-lab/summary/slow-time.txt
mysqldumpslow -s c -t 5 shared/fixtures/mysql-slow.log > ~/support-lab/summary/slow-count.txt
shared/scripts/monitor-log.sh shared/fixtures/application.log > ~/support-lab/summary/application.csv
# Expected ALERT exits 1: handle that intentionally in your wrapper.''',
['Report preserves both frequency and time rankings.','Wrapper does not discard useful ALERT output because its exit code is 1.'],
['Does one SQL statement always fit on one line?','Can identical timestamps prove causal linkage?'],
'Multiline logs require a record-aware parser. Temporal overlap narrows investigation but does not prove the app request executed a particular SQL record. Keep original evidence and clearly label derived summaries.',
'A timestamped monitoring/slow-query report and its limitations.')

lab(3,'05-scheduling','Health checks, cron and safe retention',40,
'Your script works in a terminal but fails under cron. Make paths, permissions, concurrency and output explicit.',
['Cron supplies a smaller environment and different working directory.','Use absolute paths, a controlled PATH and a lock to prevent overlapping runs.','Cleanup is dry-run first, narrowly scoped and separate from incident evidence collection.'],
['Copy scripts into ~/support-lab/bin and test from /tmp with a minimal environment.','Adapt the supplied cron example to your actual home; acquire a lock and preserve logs.','Create old dummy archives in the retention directory, run dry-run, then approve apply for those dummy files only.','Explain why replication/job checks need credentials and permissions separate from OS health checks.'],
'''mkdir -p ~/support-lab/{bin,logs,evidence}
cp shared/scripts/{health-check,retention}.sh ~/support-lab/bin/
env -i HOME="$HOME" PATH=/usr/bin:/bin "$HOME/support-lab/bin/health-check.sh"
# Review shared/templates/crontab.example; replace /home/learner01 first.
crontab -l
# Edit with crontab -e only after reviewing the exact entry.
~/support-lab/bin/retention.sh --dry-run
# --apply is only for approved dummy training archives.''',
['Scheduled execution uses absolute paths and a per-user lock.','Dry-run identifies only intended incident archives older than the retention boundary.'],
['Will a cron job inherit your interactive exported secrets?','Why can overlapping checks create a monitoring problem?'],
'Cron does not inherit the current terminal environment. Use protected credential files and least privilege. flock avoids concurrent runs; log job status and review errors. find -mtime +7 means more than seven completed 24-hour periods, not a calendar cutoff.',
'A reviewed cron entry, minimal-environment test and retention dry-run output.')

lab(3,'06-deployment-checks','Deployment verification and maintenance boundaries',20,
'A deployment health gate must catch a business failure even if the process is active.',
['Check service, readiness and one representative business response.','Timeouts and nonzero exits must propagate to the caller.','Automated diagnosis does not authorize automatic destructive repair.'],
['Extend health-check.sh to verify player 1 contains the expected synthetic result.','Run before, during and after db-auth or payment fault; choose a route matching the affected function.','State what the script may do automatically and what requires an operator decision.'],
'''shared/scripts/health-check.sh
curl --fail --silent --max-time 5 'http://127.0.0.1:8080/loyalty/player?id=1'
curl --fail --silent --max-time 5 http://127.0.0.1:8080/loyalty/payment
# Instructor selects/resets a fault; learners implement assertions, not auto-restarts.''',
['Gate fails on affected business functionality even when Tomcat remains active.','Recovery checks cover the same function that originally failed.'],
['Can a generic /health route detect every business outage?','Should the first failed probe trigger a database restart?'],
'The payment simulator can fail while database readiness remains healthy. Monitor representative business paths separately. A failed probe should emit evidence and status; restarting databases is not a safe default response.',
'A deployment-check script with documented scope and failure behavior.')

lab(3,'07-war-room','Major incident simulation',80,
'Run one of three scenarios: loyalty outage, delayed transactions, or payment integration failure. Rotate roles and deliver timed updates.',
['Separate facts, hypotheses and decisions in the incident log.','Assign owners and keep one coordinated change stream on the shared host.','Mitigation reduces impact; RCA may remain provisional at handover.'],
['Instructor assigns a scenario card and concealed fault. Use 10 minutes for briefing, 35 for investigation, 15 for mitigation/recovery and 20 for handover preparation.','Assign incident commander, investigator, scribe and communications lead; combine roles for teams of three.','Send written training updates at minutes 15, 30 and 45 using the local template, not actual external messaging.','Collect evidence, propose one change at a time, verify business recovery and document unresolved risks.'],
'''# Student tools; scenario details arrive from instructor
shared/scripts/collect-evidence.sh ~/support-lab/evidence '20 minutes ago'
shared/scripts/health-check.sh
shared/scripts/traffic.sh http://127.0.0.1:8080/loyalty 10
cp shared/templates/incident-report.md ~/support-lab/final-incident.md''',
['Timeline links symptoms, actions and results to evidence.','Updates communicate impact, current action, owner and next update time.'],
['Should you announce root cause when you only have correlation?','Can unrelated background warnings distract from the affected business function?'],
'Use instructor scenario cards. A uses db-auth; B uses latency with database decoy evidence; C uses payment while /health remains healthy. Each has distinct business recovery criteria. Grade communication of uncertainty and evidence-backed elimination of alternatives.',
'A major incident report, evidence bundle and three stakeholder updates.')

lab(3,'08-handover','Final challenges and operational handover',30,
'Another shift must continue your work without a verbal briefing.',
['A handover describes remaining work, evidence locations and safe next steps.','RCA distinguishes trigger, root cause, contributing factors and detection gaps.','Monitoring improvements need an owner and measurable acceptance criterion.'],
['Prepare a five-minute team handover using the report template.','Defend your diagnosis against a counterexample from the instructor.','Propose three monitoring/runbook improvements with owners and measurable checks.','Submit scripts, evidence index, report and recovery transcript.'],
'''# Check scripts before submission
for script in "$HOME"/support-lab/bin/*.sh; do
    [[ -f "$script" ]] || continue
    bash -n "$script" || break
done
find ~/support-lab -maxdepth 2 -type f -name '*.sh' -ls
# Review and redact all artifacts before sharing with the class.''',
['Next shift can reproduce checks and locate evidence.','Final report distinguishes confirmed facts from probable explanations.'],
['Is “monitor more” an actionable recommendation?','Can an RCA be useful when some uncertainty remains?'],
'Useful recommendations name a signal, threshold/window, owner and response. Explicit uncertainty is preferable to invented certainty. Final rubric allocates points across evidence, diagnosis, recovery, scripts and communication.',
'Operational handover and prioritized improvement backlog.')

lab(3,'09-automation-extension','Advanced automation challenge bank','Extension',
'Harden the monitoring workflow against inconvenient inputs and execution contexts.',
['Test adverse inputs, not only the happy path.','Concurrency, stale data and time zones can create false confidence.','Maintenance scripts need explicit boundaries and recovery plans.'],
['Add tests for empty files, malformed records, threshold equality, spaces in output paths and denied journal access.','Add a stale-data check and demonstrate an overlapping-run lock.','Create a dry-run cleanup test with an unrelated file that must survive.','Design replication and SQL Agent checks using the supplied evidence; describe necessary permissions and unknown states.'],
'''# Starting points for your own tests
DISK_LIMIT=100 shared/scripts/health-check.sh
shared/scripts/monitor-log.sh shared/fixtures/application.log 34 2000
# Threshold 34 suppresses rate alert, but the 2100ms request still alerts.
shared/scripts/collect-evidence.sh "$HOME/support-lab/evidence with spaces" '5 minutes ago'
shared/scripts/retention.sh --dry-run''',
['Each new behavior has a meaningful adverse-case test.','Unknown/unavailable dependencies are not reported as healthy.'],
['Can a successful parser be monitoring yesterday’s data?','What happens if a timeout exits 124 and the wrapper overwrites $?'],
'Parsing success says nothing about freshness; capture windows and last-event timestamps matter. Preserve exit codes immediately. Treat timeout, unavailable permissions and no data as explicit operational states. See reference tests for baseline expected behavior.',
'A hardening patch, test transcript and documented remaining limitations.')

MODULES[3][6]['minutes']=110
MODULES[3][6]['steps'][0]='Instructor assigns a scenario card and concealed fault. Use 10 minutes for briefing, 50 for investigation, 20 for mitigation/recovery and 30 for communications and handover preparation.'

MODULES[1][0]['commands'] += '\nhistory 10  # inspect locally; do not export commands containing secrets\nprintf "HOME=%s\\n" "$HOME"'
MODULES[1][1]['commands'] += '\nprintf "compress me\\n" > ~/support-lab/work/compression.txt\ngzip -k ~/support-lab/work/compression.txt\ngzip -dc ~/support-lab/work/compression.txt.gz'
MODULES[1][6]['steps'].append('Instructor reviews firewall indicators with ufw status verbose or nft list ruleset if installed, plus AWS security-group rules from an authorized console view. Compare this with Red Hat firewalld/SELinux evidence; never disable a firewall just to make a probe pass.')
MODULES[2][6]['steps'].append('Inspect statistics metadata using the extra query in diagnostics.sql. Treat stale statistics as a hypothesis; an approved UPDATE STATISTICS in the training database must be followed by the same workload and plan comparison, not assumed to fix every regression.')
MODULES[3][4]['steps'].append('Explain that archive retention is not a database backup. A backup policy also needs consistency, off-host storage, access controls and a tested restore; do not delete backup files merely because a log-retention rule exists.')

MODULES[3][3]['commands'] += '\nshared/scripts/monitor-pipeline.sh shared/fixtures/application.log ~/support-lab/monitoring shared/fixtures/mysql-slow.log\n# Reference wrapper preserves ALERT and reports unavailable slow-summary tooling.'

MODULES[2][0]['commands'] += '\nsystemctl status mysql --no-pager\njournalctl -u mysql --since "15 minutes ago" --utc --no-pager\n# Instructor: inspect the file identified by SHOW VARIABLES LIKE \"log_error\" when file logging is configured.'

MODULES[1][7]['steps'][0] += ' If you need a starting point, copy shared/templates/triage-starter.sh into your own bin directory and complete its TODOs.'

from revise_curriculum import apply_revision, DIAGRAMS
apply_revision(MODULES, GUIDES)

SOURCES=[
('AWS: extend an EBS filesystem','https://docs.aws.amazon.com/ebs/latest/userguide/recognize-expanded-volume-linux.html'),
('Apache Tomcat 10.1 setup','https://tomcat.apache.org/tomcat-10.1-doc/setup.html'),
('MySQL slow-query log','https://dev.mysql.com/doc/refman/8.4/en/slow-query-log.html'),
('Microsoft: install SQL Server on Ubuntu','https://learn.microsoft.com/en-us/sql/linux/quickstart-install-connect-ubuntu?view=sql-server-ver16'),
('Microsoft: active requests and permissions','https://learn.microsoft.com/en-us/sql/relational-databases/system-dynamic-management-views/sys-dm-exec-requests-transact-sql?view=sql-server-ver16'),
('GNU Bash reference','https://www.gnu.org/software/bash/manual/bash.html'),
('Red Hat: ext4 quotas','https://docs.redhat.com/en/documentation/red_hat_enterprise_linux/10/html/managing_file_systems/limiting-storage-space-usage-on-ext4-with-quotas')]

def write(path,text):
    p=ROOT/path; p.parent.mkdir(parents=True,exist_ok=True); p.write_text(text.strip()+'\n')
def bullets(items): return '\n'.join('- '+x for x in items)
def numbered(items): return '\n'.join(f'{i}. {x}' for i,x in enumerate(items,1))

def fixtures():
    import tarfile
    log='''2026-10-20T01:00:00Z level=INFO request_id=r001 path=/player status=200 latency_ms=25 error=none
2026-10-20T01:00:01Z level=ERROR request_id=r002 path=/player status=503 latency_ms=80 error=db_28000
2026-10-20T01:00:02Z level=INFO request_id=r003 path=/report status=200 latency_ms=1500 error=none
2026-10-20T01:00:03Z level=INFO request_id=r004 path=/player status=200 latency_ms=30 note=ERROR_is_a_search_term
2026-10-20T01:00:04Z level=ERROR request_id=r005 path=/payment status=502 latency_ms=2100 error=dependency_unavailable
2026-10-20T01:00:05Z level=INFO request_id=r006 path=/missing status=404 latency_ms=5 error=none
malformed training record without required fields
'''
    write(Path('shared/fixtures/application.log'),log)
    write(Path('shared/fixtures/latency.csv'),'request_id,latency_ms\nr001,25\nr002,80\nr003,1500\nr004,30\nr005,2100\nr006,5')
    for name in ['normal.txt','status report.txt','-dangerous-name.txt','.hidden.txt','line\nbreak.txt']:
        write(Path('shared/fixtures/filenames')/name,'Synthetic training data; filename is intentionally awkward.')
    with tarfile.open(ROOT/'shared/fixtures/support-fixtures.tar.gz','w:gz') as t:
        for name in ['application.log','latency.csv','filenames']:
            t.add(ROOT/'shared/fixtures'/name,arcname=name)
    slow=[]
    for i,(seconds,query) in enumerate([(1.2,'SELECT SUM(amount) FROM transactions WHERE player_id=42;'),(1.3,'SELECT SUM(amount) FROM transactions WHERE player_id=43;'),(1.1,'SELECT SUM(amount) FROM transactions WHERE player_id=44;'),(3.0,'SELECT SLEEP(3);')]):
        slow.append(f'# Time: 2026-10-20T01:00:0{i}.000000Z\n# User@Host: training[training] @ localhost [] Id: {i+10}\n# Query_time: {seconds:.6f} Lock_time: 0.000100 Rows_sent: 1 Rows_examined: {100000 if i<3 else 0}\nuse support_lab;\nSET timestamp=179245800{i};\n{query}\n')
    write(Path('shared/fixtures/mysql-slow.log'),'\n'.join(slow))
    write(Path('shared/fixtures/replication.json'),json.dumps({'provenance':'SYNTHETIC teaching fixture; not captured from EC2','captured_at':'2026-10-20T01:05:00Z','primary_last_write':'2026-10-20T01:04:50Z','report_route':'replica','replica_last_visible_write':'2026-10-20T00:58:00Z','Replica_IO_Running':'Yes','Replica_SQL_Running':'No','Seconds_Behind_Source':None,'Last_SQL_Errno':1062,'Last_SQL_Error':'Duplicate entry for synthetic training key','Relay_Log_File':'replica-relay-bin.000012'},indent=2))
    write(Path('shared/fixtures/sqlserver-waits.csv'),'captured_at,wait_type,wait_time_ms\n2026-10-20T01:00:00Z,LCK_M_X,100000\n2026-10-20T01:00:00Z,SOS_SCHEDULER_YIELD,50000\n2026-10-20T01:00:00Z,PAGEIOLATCH_SH,80000\n2026-10-20T01:01:00Z,LCK_M_X,112000\n2026-10-20T01:01:00Z,SOS_SCHEDULER_YIELD,51000\n2026-10-20T01:01:00Z,PAGEIOLATCH_SH,80200')
    write(Path('shared/fixtures/query-store.csv'),'interval_utc,query_id,plan_id,count_executions,avg_duration_us,avg_cpu_us,avg_logical_reads\n2026-10-20T00:00:00Z,17,3,100,50000,20000,200\n2026-10-20T01:00:00Z,17,8,100,800000,500000,20000')
    write(Path('shared/fixtures/job-history.csv'),'job,step_id,run_status,message\nReportRefresh,1,0,SELECT permission denied on synthetic reporting table\nReportRefresh,0,0,The job failed at step 1')
    write(Path('shared/fixtures/correlation.csv'),'request_id,start_utc,pool_wait_ms,sql_ms,other_ms,http_total_ms,attempts,operation\nr900,2026-10-20T01:00:00Z,4000,40,960,5000,3,player_lookup\nr901,2026-10-20T01:00:02Z,0,35,25,60,1,player_lookup')
    write(Path('shared/fixtures/README.md'),'''# Synthetic evidence fixtures
All records in this folder are authored teaching data, not logs collected from customers or the EC2 server. Times are UTC. Do not infer that separate files represent one actual incident unless an exercise explicitly joins them.

- application.log: six valid requests and one malformed row; HTTP status and latency are explicit fields.
- mysql-slow.log: four query executions; mysqldumpslow groups the three similar report queries.
- replication.json: stopped SQL applier despite a running receiver.
- sqlserver-waits.csv: two cumulative captures, 60 seconds apart; compute deltas.
- query-store.csv: two comparable synthetic intervals; duration/CPU are microseconds.
- job-history.csv: a failed step and aggregate job result; a different failure from the live intentional-THROW job.
- correlation.csv: synthetic pooled application, unlike the simple live app.
- platform-incidents.md: DNS/TLS/SELinux/IIS and high-CPU evidence cards.
- support-fixtures.tar.gz: starter files for Linux command practice; no executable payloads.
''')

def documents():
    titles={1:'Linux Fundamentals and Incident Triage',2:'MySQL and SQL Server Diagnostics',3:'Automation and Major Incident Simulation'}
    for m,labs in MODULES.items():
        agenda='\n'.join(f"| [{l['title']}](exercises/{l['slug']}.md) | {l['minutes']} |" for l in labs)
        assert sum(l['minutes'] for l in labs if isinstance(l['minutes'],int))==370
        write(Path(f'module{m}/README.md'),f'''# Module {m}: {titles[m]}

One training day: 390 instructional minutes, 60-minute lunch and two 15-minute breaks. Advanced extensions are optional. Teams have 3–4 learners; rotate roles and schedule shared service changes.

[PowerPoint deck](Module{m}.pptx) · [Learner setup](../shared/learner-start.md) · [Assessment rubric](../shared/assessment-rubric.md)

## Agenda
| Exercise | Instructional minutes |
|---|---:|
{agenda}
| Checkpoints and recovery reserve | 20 |

The exercise blocks total 370 minutes. Use two 10-minute checkpoints to review evidence, recover interrupted work and confirm readiness. Demonstrations and extensions are labeled in each exercise; extensions do not displace the final assessment.

## Before starting
Read the [learner terminal/setup guide](../shared/learner-start.md). Each exercise now includes numbered actions, execution context, expected results and specific cleanup. Final challenges retain independent diagnosis with optional hints.

Use the course repository root for commands unless the exercise changes directory. Copy evidence into `~/support-lab`, never edit original fixtures. Commands marked instructor require the instructor; students do not need blanket sudo. Use individual SSH credentials. Modules 2 and 3 build on the prior day's application and evidence skills.

## Completion
Submit the deliverables listed in each required exercise. Keep passwords, private keys and unreviewed sensitive logs out of submissions. Distinguish synthetic evidence from live observations. The instructor scores evidence, diagnosis, recovery, automation and communication using the shared rubric.
''')
        answers=[]
        for i,l in enumerate(labs):
            write(Path(f'module{m}/exercises/{l["slug"]}.md'), render_exercise(m,i+1,l))
            answers.append(f"## {l['title']}\n\nSlides {5+i*5}–{9+i*5}. Time: {l['minutes']}.\n\n**Expected reasoning:** {l['answer']}\n\n**Hint ladder:** First ask which layer the evidence describes. Next ask for a discriminating command from the worksheet. Finally point to the relevant fixture or reference script; do not reveal the hidden fault immediately.\n\n**Acceptance:** {' '.join(l['checks'])}\n\n**Submission:** {l['output']}\n")
            g=GUIDES[m,i+1]
            answers[-1] += "\n**Guided walkthrough checkpoints:**\n\n" + "\n".join(f"{j}. {step['title']} — {step['expected']}" for j,step in enumerate(g['steps'],1)) + "\n\n**Exercise-specific recovery:** " + g['cleanup'] + "\n" + '\n**Instructor execution sheet:**\n\n' + '\n\n'.join(s['title']+'\n\n'+s['action']+'\n\n```'+s['language']+'\n'+s['code']+'\n```' for s in g['steps'] if s['where'].startswith('Instructor') and s['code'])
        write(Path(f'module{m}/instructor/guide.md'),f"# Module {m} instructor guide\n\nKeep this answer key separate from student handouts. Rehearse setup and fault reset before class. Speaker notes include the same reasoning for live delivery. Review [shared setup](../../shared/instructor/setup-and-delivery.md), [scenario cards/rubric](../../shared/instructor/scenarios.md) and [storage procedures](../../shared/instructor/storage.md).\n\n"+'\n'.join(answers))
    write(Path('shared/SOURCES.md'),'# Technical references\n\nConsulted during authoring. Examples are original training material; fixtures are synthetic. Recheck installation/version-sensitive instructions before delivery.\n\n'+ '\n'.join(f'- [{name}]({url})' for name,url in SOURCES))
    # Human-reviewable, editable content source for deck rebuilding.
    write(Path('tools/course-content.json'),json.dumps(MODULES,indent=2))

def decks():
    import pymupdf
    from PIL import Image
    from collections import Counter
    from pptx import Presentation
    from pptx.util import Inches,Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
    assets=ROOT/'shared/branding'; assets.mkdir(exist_ok=True)
    pdf=pymupdf.open(ROOT/'Application Support and Monitoring Course Outline 3.pdf')
    # Extract logo region from the supplied outline, without recreating the trademark.
    pix=pdf[0].get_pixmap(matrix=pymupdf.Matrix(3,3),clip=pymupdf.Rect(288,40,568,88))
    logo=assets/'trainocate-logo.png'; pix.save(logo)
    im=Image.open(logo).convert('RGB'); counts=Counter(im.getdata())
    orange=next(c for c,n in counts.most_common() if c[0]>200 and 40<c[1]<140 and c[2]<90)
    gray=(84,84,84); ink=(28,32,36); muted=(102,110,118); white=(255,255,255)
    write(Path('shared/branding/README.md'),f'# Branding\n\nLogo cropped from the supplied outline. Sampled orange: #{orange[0]:02X}{orange[1]:02X}{orange[2]:02X}. Charcoal: #545454. White background, dark body text. Logo is retained for the requested Trainocate training materials.\n')
    titles={1:'Linux fundamentals\nand incident triage',2:'Database performance\nand incident diagnostics',3:'Automation, monitoring\nand major incidents'}
    def box(s,x,y,w,h,text,size=24,color=ink,bold=False,fill=None,font='Aptos'):
        if fill:
            sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor(*fill);sh.line.fill.background()
        else: sh=s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf=sh.text_frame; tf.word_wrap=True; tf.vertical_anchor=MSO_ANCHOR.TOP
        tf.margin_left=Inches(.13);tf.margin_right=Inches(.13);tf.margin_top=Inches(.08);tf.margin_bottom=Inches(.04)
        for j,line in enumerate(text.split('\n')):
            p=tf.paragraphs[0] if j==0 else tf.add_paragraph();p.text=line;p.alignment=PP_ALIGN.LEFT;p.font.name=font;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor(*color);p.space_after=Pt(8 if size<=24 else 12)
        return sh
    for m,labs in MODULES.items():
        prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
        def slide(title,tag='APPLICATION SUPPORT AND MONITORING'):
            s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor(*white)
            box(s,0,0,.16,7.5,'',fill=orange)
            box(s,.48,.2,9,.5,tag,9 if len(tag)>60 else 11,muted,True)
            s.shapes.add_picture(str(logo), Inches(10.28), Inches(.2),width=Inches(2.55))
            box(s,.5,.88,12.2,1.15,title,30,ink,True)
            box(s,.55,7.08,11,.24,f'TRAINOCATE  •  MODULE {m}  •  SYNTHETIC TRAINING ENVIRONMENT',9,muted)
            box(s,12.15,7.04,.6,.3,str(len(prs.slides)),10,muted)
            return s
        s=slide(f'Module {m}', 'THREE-DAY PRACTICAL TRAINING')
        box(s,.6,2.05,11.8,2.2,titles[m],40,white,True,gray)
        box(s,.7,4.7,11.8,1.3,'One day • 390 instructional minutes\nInvestigate → explain → mitigate → verify → hand over',23)
        s.notes_slide.notes_text_frame.text='Begin with learner experience and expected outputs. The host is shared: only one team injects a fault at a time. The supplied PDF is the curriculum authority; Ubuntu is the live platform.'
        s=slide('Today’s route')
        ag=[f"{l['minutes']:>3} min  |  {l['title']}" for l in labs if isinstance(l['minutes'],int)] + ['20 min  |  Checkpoints and recovery reserve']
        box(s,.65,2.02,12,4.8,'\n'.join(ag),17)
        s.notes_slide.notes_text_frame.text='390 minutes plus 60-minute lunch and two 15-minute breaks. Advanced challenge banks are extension work. Split blocks around breaks as needed; do not silently remove the capstone.'
        s=slide('The lab system and its boundaries')
        for x,w,t in [(.7,3.7,'LEARNER\nSSH + individual account\nTeam workspace'),(4.8,3.7,'APPLICATION HOST\nTomcat + MySQL\nLoyalty training app'),(8.9,3.7,'DATABASE LAB\nSQL Server 2022\nQuery Store + Agent')]: box(s,x,2.1,w,1.9,t,21,white,True,gray)
        box(s,.8,4.5,11.7,1.65,'Request IDs connect application evidence.\nSynthetic packs cover replication, IIS, SELinux and pooled-app incidents.\nFault injection and storage administration are instructor-controlled.',22)
        s.notes_slide.notes_text_frame.text='See shared/instructor/setup-and-delivery.md. Current EC2 capacity is not the classroom recommendation. Keep ports private; use SSH forwarding. The simple app does not implement a connection pool or real payments.'
        s=slide('How to work an incident')
        box(s,.7,2.05,11.8,4.6,'1  State impact and capture UTC time.\n2  Preserve evidence before changing the system.\n3  Test a hypothesis with a discriminating check.\n4  Make one scoped change with a rollback.\n5  Verify the affected business function and communicate.',25)
        s.notes_slide.notes_text_frame.text='Assign investigator, operator, scribe/reviewer. Use the shared report template. Empty output, missing permissions and unknown states are not proof of health.'
        for n,l in enumerate(labs,1):
            g=GUIDES[m,n]
            notes=f"Exercise: module{m}/exercises/{l['slug']}.md\nDuration: {l['minutes']}\n\n"+'\n'.join(l['steps'])+'\n\nINSTRUCTOR ANSWER: '+l['answer']+'\n\nDeliverable: '+l['output']
            s=slide(l['title'],f"LEARN  •  {l['minutes']} {'MINUTES' if isinstance(l['minutes'],int) else ''}")
            box(s,.65,2.0,12,1.25,l['scenario'],23,white,False,gray)
            box(s,.7,3.55,11.9,2.95,'\n'.join('• '+v for v in l['concepts']),20)
            s.notes_slide.notes_text_frame.text=notes
            diagram=DIAGRAMS.get((m,n))
            if diagram:
                title, stages, explanation=diagram
                s=slide(title,l['title'].upper())
                for j,label in enumerate(stages):
                    box(s,.7+j*4.15,2.2,3.55,1.65,label,23,white,True,gray)
                    if j<2: box(s,4.27+j*4.15,2.65,.5,.6,'→',25,orange,True)
                box(s,.8,4.5,11.7,1.6,explanation,24)
            else:
                s=slide('Your investigation sequence',l['title'].upper())
                steps=[v['title'] for v in g['steps'] if not v['title'].startswith('Live extension')][:6]
                box(s,.7,2.0,11.9,4.65,'\n'.join(f'{i}  {v}' for i,v in enumerate(steps,1)),22)
            s.notes_slide.notes_text_frame.text=notes
            chosen=g['slide_step']
            s=slide(chosen['title'] if chosen else 'Choose a discriminating check',l['title'].upper())
            snippet=chosen['code'] if chosen else 'Use the numbered worksheet.\nExplain which hypothesis your next check tests.'
            box(s,.65,2.0,12,3.65,snippet,16,white,False,gray,'DejaVu Sans Mono')
            box(s,.75,5.9,11.6,.75,'Read the complete worksheet before executing. Instructor-only changes stay supervised.',18,muted)
            s.notes_slide.notes_text_frame.text=notes+'\n\nFULL WORKSHEET COMMANDS:\n'+l['commands']
            s=slide('Prove it — and challenge the obvious answer',l['title'].upper())
            box(s,.7,2.0,11.9,1.6,'\n'.join('✓ '+v for v in l['checks']),21)
            box(s,.7,3.9,11.9,2.6,'\n'.join('? '+v for v in l['tricks']),23,ink,False,(246,246,246))
            s.notes_slide.notes_text_frame.text=notes
            s=slide('Debrief and deliverable',l['title'].upper())
            box(s,.7,2.05,11.9,1.5,l['output'],27,white,True,gray)
            box(s,.75,3.95,11.7,2.2,'What did the evidence prove?\nWhat alternative did you rule out?\nHow did you verify recovery?\nWhat remains uncertain?',25)
            s.notes_slide.notes_text_frame.text=notes
        s=slide('Reference shelf and completion checks')
        box(s,.7,2.0,11.9,3.3,'Use shared/SOURCES.md for official technical references.\nSubmit scripts, evidence, interpretations and recovery checks.\nKeep private keys and passwords out of submissions.\nComplete the core labs before attempting extension challenges.',24)
        box(s,.7,5.4,11.9,.8,'Evidence 25 • Diagnosis 25 • Recovery 20 • Automation 20 • Communication 10',19,white,True,gray)
        s.notes_slide.notes_text_frame.text='\n'.join(f'{n}: {u}' for n,u in SOURCES)+'\nPass: 70/100 with required evidence, script/outline and recovery verification. See rubric.'
        prs.save(ROOT/f'module{m}/Module{m}.pptx')
    write(Path('tools/build-requirements.txt'),'python-pptx==1.0.2\nPyMuPDF>=1.26,<2\nPillow>=11,<13')

if __name__=='__main__':
    import argparse
    parser=argparse.ArgumentParser()
    parser.add_argument('--documents-only', action='store_true', help='Refresh guided Markdown without changing decks or fixtures')
    args=parser.parse_args()
    if args.documents_only:
        documents()
        print('Updated 27 guided exercises and instructor checkpoints; decks and fixtures unchanged.')
    else:
        fixtures(); documents(); decks()
        print('Created guided exercises, answer keys, fixtures and three PowerPoint decks.')
