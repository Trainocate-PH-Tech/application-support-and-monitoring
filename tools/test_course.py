"""Local checks; no remote hosts, privileged operations or live fault injection."""
import csv, io, os, re, subprocess, tempfile, unittest
from pathlib import Path
from pptx import Presentation
ROOT=Path(__file__).resolve().parents[1]
SCRIPTS=ROOT/'shared/scripts'
def run(script,*args,env=None):
    return subprocess.run(['bash',str(SCRIPTS/script),*map(str,args)],capture_output=True,text=True,env=env)
class CourseTests(unittest.TestCase):
    def test_shell_syntax(self):
        for p in SCRIPTS.glob('*.sh'):
            r=subprocess.run(['bash','-n',str(p)],capture_output=True,text=True)
            self.assertEqual(r.returncode,0,(p,r.stderr))
    def test_fixture_alert(self):
        r=run('monitor-log.sh',ROOT/'shared/fixtures/application.log')
        self.assertEqual(r.returncode,1)
        row=list(csv.DictReader(io.StringIO(r.stdout)))[0]
        self.assertEqual(row,dict(requests='6',errors='2',error_percent='33.33',slow_requests='2',malformed='1',state='ALERT'))
    def test_parser_boundaries(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'log with spaces'
            for text,code in [('',2),('bad record\n',2),('status=200 latency_ms=10\n',0),('status=500 latency_ms=1000\n',1)]:
                p.write_text(text);self.assertEqual(run('monitor-log.sh',p,100,1000).returncode,code)
            self.assertEqual(run('monitor-log.sh',p,'invalid').returncode,2)
    def test_monitor_pipeline_keeps_alert_output(self):
        with tempfile.TemporaryDirectory() as d:
            r=run('monitor-pipeline.sh',ROOT/'shared/fixtures/application.log',d)
            self.assertEqual(r.returncode,1,r.stderr)
            folder=next(Path(d).glob('window-*'))
            self.assertIn('ALERT',(folder/'application.csv').read_text())
            self.assertIn('application_exit=1',(folder/'manifest.txt').read_text())
    def test_retention_scope(self):
        with tempfile.TemporaryDirectory() as d:
            base=Path(d)/'support-lab/evidence';base.mkdir(parents=True)
            old=base/'incident-old.tar.gz';other=base/'unrelated.tar.gz';new=base/'incident-new.tar.gz'
            for p in (old,other,new):p.write_text('dummy')
            os.utime(old,(1,1));os.utime(other,(1,1))
            env={**os.environ,'HOME':d}
            self.assertEqual(run('retention.sh',env=env).returncode,0);self.assertTrue(old.exists())
            self.assertEqual(run('retention.sh','--apply',env=env).returncode,0)
            self.assertFalse(old.exists());self.assertTrue(other.exists());self.assertTrue(new.exists())
    def test_health_failures_with_stubs(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)
            def stub(name,code):
                f=p/name;f.write_text('#!/bin/sh\n'+code+'\n');f.chmod(0o755)
            stub('systemctl','exit "${MOCK_SERVICE_RC:-0}"')
            stub('curl','printf "%s" "${MOCK_HTTP:-200}"; exit "${MOCK_CURL_RC:-0}"')
            stub('df','printf "Filesystem 1024-blocks Used Available Capacity Mounted\\n/dev/fake 100 30 70 ${MOCK_DISK:-30}%% /\\n"')
            env={**os.environ,'PATH':d+':'+os.environ['PATH']}
            self.assertEqual(run('health-check.sh',env=env).returncode,0)
            for change in ({'MOCK_SERVICE_RC':'3'},{'MOCK_HTTP':'503'},{'MOCK_CURL_RC':'28'},{'MOCK_DISK':'85'}):
                self.assertEqual(run('health-check.sh',env={**env,**change}).returncode,1)
            self.assertEqual(run('health-check.sh',env={**env,'DISK_LIMIT':'bad'}).returncode,2)
    def test_partial_evidence_and_masking(self):
        with tempfile.TemporaryDirectory() as d:
            bin=Path(d)/'bin';bin.mkdir()
            for name,code in [('systemctl','echo service_failed; exit 3'),('journalctl','echo "password=syntheticsecret token=synthetictoken"')]:
                p=bin/name;p.write_text('#!/bin/sh\n'+code+'\n');p.chmod(0o755)
            env={**os.environ,'PATH':str(bin)+':'+os.environ['PATH']}
            r=run('collect-evidence.sh',Path(d)/'output with spaces',env=env)
            self.assertEqual(r.returncode,1,r.stderr)
            folder=next((Path(d)/'output with spaces').glob('incident-*/'))
            self.assertIn('service\t3',(folder/'manifest.tsv').read_text())
            text=(folder/'journal.txt').read_text();self.assertNotIn('syntheticsecret',text);self.assertIn('[REDACTED]',text)
            self.assertTrue(Path(str(folder)+'.tar.gz').exists())
            self.assertEqual(folder.stat().st_mode & 0o777,0o700)
    def test_decks_and_exercises(self):
        for m in (1,2,3):
            prs=Presentation(ROOT/f'module{m}/Module{m}.pptx')
            self.assertEqual(len(prs.slides),50)
            self.assertEqual(len(list((ROOT/f'module{m}/exercises').glob('*.md'))),9)
            for slide in prs.slides:
                self.assertTrue(slide.notes_slide.notes_text_frame.text.strip())
                for s in slide.shapes:
                    self.assertGreaterEqual(s.left,0);self.assertGreaterEqual(s.top,0)
                    self.assertLessEqual(s.left+s.width,prs.slide_width+10)
                    self.assertLessEqual(s.top+s.height,prs.slide_height+10)
    def test_markdown_local_links(self):
        for p in ROOT.rglob('*.md'):
            if '.build' in p.parts or 'target' in p.parts:continue
            for target in re.findall(r'\]\(([^)]+)\)',p.read_text()):
                if target.startswith(('https://','http://','#')):continue
                self.assertTrue((p.parent/target.split('#')[0]).exists(),f'{p}: {target}')
if __name__=='__main__':unittest.main(verbosity=2)
