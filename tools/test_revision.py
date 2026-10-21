"""Regression checks for revised teaching flow and audience-specific artifacts."""
import io
import os
from pathlib import Path, PurePosixPath
import posixpath
import re
import subprocess
import tempfile
import unittest
import zipfile
from pptx import Presentation
import pymupdf
import build_course
from revise_curriculum import HEALTH, FOUNDATION, SYNTAX_CHECK

ROOT=Path(__file__).resolve().parents[1]

class RevisionTests(unittest.TestCase):
    def test_pacing_and_prerequisite(self):
        for m,labs in build_course.MODULES.items():
            self.assertEqual(sum(l['minutes'] for l in labs if isinstance(l['minutes'],int))+20,390)
            self.assertIn('Checkpoints and recovery reserve',(ROOT/f'module{m}/README.md').read_text())
        first=build_course.GUIDES[1,4]
        self.assertTrue(any(s['code']==FOUNDATION for s in first['steps']))
        self.assertIn('Reuse triage.sh',build_course.GUIDES[1,8]['steps'][1]['action'])

    def test_health_scaffold_adverse_cases(self):
        with tempfile.TemporaryDirectory() as td:
            p=Path(td); script=p/'health.sh';script.write_text(HEALTH)
            bin=p/'bin';bin.mkdir()
            for name,body in [('systemctl','exit "${SERVICE_RC:-0}"'),('curl','exit "${CURL_RC:-0}"')]:
                f=bin/name;f.write_text('#!/bin/sh\n'+body+'\n');f.chmod(0o755)
            env={**os.environ,'PATH':str(bin)+':'+os.environ['PATH']}
            def run(args=(),change=None):
                return subprocess.run(['/bin/bash',str(script),*args],env={**env,**(change or {})},capture_output=True,text=True,timeout=10).returncode
            self.assertEqual(run(),0)
            self.assertEqual(run(change={'CURL_RC':'7'}),1)
            self.assertEqual(run(change={'SERVICE_RC':'3'}),1)
            self.assertEqual(run(('support-tomcat','http://localhost','extra')),2)
            self.assertEqual(run(('support-tomcat','file:///etc/passwd')),2)
            self.assertEqual(run(('','http://localhost')),2)
            self.assertEqual(run(change={'PATH':str(p/'empty')}),2)
            # A URL with spaces is one array argument; no shell evaluation or word splitting.
            (bin/'curl').write_text('#!/bin/sh\n[ "$#" -eq 6 ] && [ "$6" = "http://localhost/a b" ]\n')
            self.assertEqual(run(('support-tomcat','http://localhost/a b')),0)

    def test_foundation_repeat_preserves_script(self):
        with tempfile.TemporaryDirectory() as td:
            p=Path(td); (p/'support-lab/evidence/m1-04').mkdir(parents=True)
            env={**os.environ,'HOME':td}
            subprocess.run(['bash','-c',FOUNDATION],env=env,capture_output=True,check=True)
            script=p/'support-lab/bin/triage.sh'
            script.write_text(script.read_text()+'\n# learner-owned marker\n')
            subprocess.run(['bash','-c',FOUNDATION],env=env,capture_output=True,check=True)
            self.assertIn('learner-owned marker',script.read_text())

    def test_syntax_loop_checks_second_script(self):
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'support-lab/bin';p.mkdir(parents=True)
            (p/'a.sh').write_text('echo valid\n')
            (p/'b.sh').write_text('if then\n')
            run=subprocess.run(['bash','-c',SYNTAX_CHECK],env={**os.environ,'HOME':td},capture_output=True,text=True)
            self.assertIn('syntax_exit=1',run.stdout)
            (p/'b.sh').write_text('echo valid\n')
            run=subprocess.run(['bash','-c',SYNTAX_CHECK],env={**os.environ,'HOME':td},capture_output=True,text=True)
            self.assertIn('syntax_exit=0',run.stdout)

    def test_slide_snippets_match_walkthroughs(self):
        for m,labs in build_course.MODULES.items():
            prs=Presentation(ROOT/f'module{m}/Module{m}.pptx')
            for n,_ in enumerate(labs,1):
                chosen=build_course.GUIDES[m,n]['slide_step']
                if chosen:
                    slide=prs.slides[4+(n-1)*5+2]
                    texts=[s.text for s in slide.shapes if s.has_text_frame]
                    self.assertIn(chosen['code'],texts)
                    worksheet=(ROOT/f'module{m}/exercises/{labs[n-1]["slug"]}.md').read_text()
                    self.assertIn(chosen['code'],worksheet)

    def test_student_archive_boundaries(self):
        with zipfile.ZipFile(ROOT/'Application-Support-Student.zip') as z:
            names=z.namelist()
            self.assertIsNone(z.testzip())
            self.assertEqual(len([n for n in names if '/exercises/' in n]),27)
            for n in names:
                self.assertNotIn('/instructor/',n)
                self.assertFalse(n.startswith('tools/'))
                self.assertFalse(n.endswith(('fault.sh','setup-ubuntu.sh','storage-lab.sh','.pem','.key')))
                if n.startswith('shared/scripts/') and n.endswith('.sh'):
                    self.assertTrue((z.getinfo(n).external_attr >> 16) & 0o111,n)
                if n.endswith('.pptx'):
                    prs=Presentation(io.BytesIO(z.read(n)))
                    for slide in prs.slides:
                        self.assertEqual(slide.notes_slide.notes_text_frame.text,'Learner edition. Follow the numbered exercise and record your own evidence.')
                if n.endswith('.md'):
                    text=z.read(n).decode()
                    self.assertNotIn('R18 wrapped the indexed',text)
                    self.assertNotIn('02:12 UTC',text)
                    for target in re.findall(r'\]\(([^)]+)\)',text):
                        if target.startswith(('http://','https://','#')): continue
                        resolved=posixpath.normpath(str(PurePosixPath(n).parent/target.split('#')[0]))
                        self.assertIn(resolved,names,(n,target))
            for n in ['shared/sql/mysql/diagnostics.sql','shared/sql/sqlserver/diagnostics.sql','shared/sql/sqlserver/workload.sql']:
                self.assertIn(n,names)

    def test_instructor_package_retains_teaching_tools(self):
        with zipfile.ZipFile(ROOT/'Application-Support-Training.zip') as z:
            for n in ['shared/instructor/scenarios.md','shared/instructor/platform-answers.md','tools/revise_curriculum.py','shared/scripts/fault.sh','shared/app/releases/loyalty.war']:
                self.assertIn(n,z.namelist())
            self.assertIsNone(z.testzip())

    def test_pdf_pages_and_text_bounds(self):
        for m in (1,2,3):
            doc=pymupdf.open(ROOT/f'module{m}/Module{m}.pdf')
            self.assertEqual(len(doc),50)
            for i,page in enumerate(doc):
                for block in page.get_text('dict')['blocks']:
                    for line in block.get('lines',[]):
                        for span in line['spans']:
                            r=pymupdf.Rect(span['bbox'])
                            self.assertTrue(page.rect.contains(r),(m,i+1,span['text']))
                self.assertNotIn('INSTRUCTOR ANSWER',page.get_text())

if __name__=='__main__': unittest.main(verbosity=2)
